import csv
import os

from itertools import islice
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".xlsx",
}


class SourceDataError(RuntimeError):
    pass


class FileDataSource:
    """
    Read-only access to external CSV/XLSX files.

    The source files remain outside ASTRA.
    Nothing is copied into PostgreSQL.
    """

    def __init__(
        self,
        root_path: str | Path | None = None,
    ):
        configured_path = (
            root_path
            or os.getenv(
                "SOURCE_DATA_PATH",
                "/data/source",
            )
        )

        self.root_path = Path(
            configured_path
        ).expanduser().resolve()

    def exists(self) -> bool:
        return self.root_path.exists()

    def list_files(self) -> list[Path]:
        if not self.root_path.exists():
            raise SourceDataError(
                f"Source path does not exist: "
                f"{self.root_path}"
            )

        if self.root_path.is_file():
            self._validate_extension(
                self.root_path
            )

            return [
                self.root_path
            ]

        files = [
            path
            for path in self.root_path.rglob("*")
            if (
                path.is_file()
                and path.suffix.lower()
                in SUPPORTED_EXTENSIONS
            )
        ]

        return sorted(files)

    def preview(
        self,
        file_name: str | None = None,
        limit: int = 5,
    ) -> list[dict[str, Any]]:
        if limit < 1 or limit > 100:
            raise ValueError(
                "limit must be between 1 and 100"
            )

        path = self._resolve_file(
            file_name
        )

        suffix = path.suffix.lower()

        if suffix == ".csv":
            return self._preview_csv(
                path,
                limit,
            )

        if suffix == ".xlsx":
            return self._preview_xlsx(
                path,
                limit,
            )

        raise SourceDataError(
            f"Unsupported file type: {suffix}"
        )

    def _resolve_file(
        self,
        file_name: str | None,
    ) -> Path:
        files = self.list_files()

        if file_name is None:
            if len(files) == 1:
                return files[0]

            if not files:
                raise SourceDataError(
                    "No source files found"
                )

            raise SourceDataError(
                "Multiple source files found. "
                "Specify file_name."
            )

        if self.root_path.is_file():
            candidate = self.root_path
        else:
            candidate = (
                self.root_path
                / file_name
            ).resolve()

            try:
                candidate.relative_to(
                    self.root_path
                )
            except ValueError as exc:
                raise SourceDataError(
                    "File is outside SOURCE_DATA_PATH"
                ) from exc

        if not candidate.exists():
            raise SourceDataError(
                f"Source file not found: "
                f"{candidate}"
            )

        if not candidate.is_file():
            raise SourceDataError(
                f"Not a file: {candidate}"
            )

        self._validate_extension(
            candidate
        )

        return candidate

    @staticmethod
    def _validate_extension(
        path: Path,
    ) -> None:
        if (
            path.suffix.lower()
            not in SUPPORTED_EXTENSIONS
        ):
            raise SourceDataError(
                "Only CSV and XLSX "
                "sources are supported"
            )

    @staticmethod
    def _preview_csv(
        path: Path,
        limit: int,
    ) -> list[dict[str, Any]]:
        with path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as file:
            sample = file.read(
                64 * 1024
            )

            file.seek(0)

            try:
                dialect = csv.Sniffer().sniff(
                    sample,
                    delimiters=",;\t|",
                )
            except csv.Error:
                dialect = csv.excel

            reader = csv.DictReader(
                file,
                dialect=dialect,
            )

            return [
                dict(row)
                for row in islice(
                    reader,
                    limit,
                )
            ]

    @staticmethod
    def _preview_xlsx(
        path: Path,
        limit: int,
    ) -> list[dict[str, Any]]:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.active

            rows = worksheet.iter_rows(
                values_only=True
            )

            headers_row = next(
                rows,
                None,
            )

            if headers_row is None:
                return []

            headers = [
                str(value)
                if value is not None
                else f"column_{index + 1}"
                for index, value
                in enumerate(headers_row)
            ]

            result: list[
                dict[str, Any]
            ] = []

            for row in islice(
                rows,
                limit,
            ):
                result.append(
                    dict(
                        zip(
                            headers,
                            row,
                        )
                    )
                )

            return result

        finally:
            workbook.close()
